import io
import os
import uuid

import httpx
import pytest
from cryptography.exceptions import InvalidTag
from docx import Document

from app.core.crypto import blind_index, decrypt_field, encrypt_field
from app.services import backup_service as bs
from app.services import line_service
from app.services.export_service import ExportError, render_docx


def test_field_encryption_roundtrip_and_context_binding():
    blob = encrypt_field("1234567890121", "personnel.citizen_id")
    assert b"1234567890121" not in blob
    assert decrypt_field(blob, "personnel.citizen_id") == "1234567890121"
    with pytest.raises(InvalidTag):   # ย้าย ciphertext ไปคอลัมน์อื่น → ถอดไม่ได้
        decrypt_field(blob, "vendors.bank_account")
    assert blind_index("1234567890121") == blind_index("1234567890121") != blind_index("1234567890122")


def _roundtrip(data: bytes) -> bytes:
    enc = io.BytesIO()
    bs.encrypt_stream(io.BytesIO(data), enc)
    out = io.BytesIO()
    bs.decrypt_stream(io.BytesIO(enc.getvalue()), out)
    return out.getvalue()


def test_backup_encryption_roundtrip(monkeypatch):
    monkeypatch.setattr(bs, "CHUNK", 1000)   # บังคับให้มีหลาย chunk
    for size in (0, 1, 999, 1000, 1001, 5555):
        data = os.urandom(size)
        assert _roundtrip(data) == data


def test_backup_encryption_detects_tamper_and_truncation(monkeypatch):
    monkeypatch.setattr(bs, "CHUNK", 1000)
    enc = io.BytesIO()
    bs.encrypt_stream(io.BytesIO(os.urandom(3500)), enc)
    raw = enc.getvalue()

    flipped = bytearray(raw)
    flipped[200] ^= 1
    with pytest.raises(bs.BackupError):
        bs.decrypt_stream(io.BytesIO(bytes(flipped)), io.BytesIO())

    first_chunk_end = len(bs.MAGIC) + 8 + 5 + 1000 + 16
    with pytest.raises(bs.BackupError, match="ตัดท้าย"):
        bs.decrypt_stream(io.BytesIO(raw[:first_chunk_end]), io.BytesIO())
    with pytest.raises(bs.BackupError):
        bs.decrypt_stream(io.BytesIO(raw + b"x"), io.BytesIO())


def test_backup_rejects_path_traversal():
    for name in ("../../etc/passwd", "procurement-20260101T000000.dump.enc/../x", "evil.dump.enc"):
        with pytest.raises(bs.BackupError):
            bs._safe_backup_path(name)


def test_line_push_retries_then_succeeds(monkeypatch):
    monkeypatch.setattr(line_service.time, "sleep", lambda _: None)
    calls = []

    def handler(req: httpx.Request) -> httpx.Response:
        calls.append(req)
        return httpx.Response(500 if len(calls) == 1 else 200, json={})

    client = httpx.Client(base_url="https://api.line.me", transport=httpx.MockTransport(handler))
    key = uuid.uuid4()
    msg = line_service.build_status_message(
        project_code="P-001", project_title="<script>alert(1)</script>", method_name="วิธีเฉพาะเจาะจง",
        step_name="บันทึกตกลงราคา", status="SUBMITTED", budget_text="80,000.00 บาท", actor_name="นายทดสอบ")
    assert line_service.push_message("Cgroup", msg, key, client=client) == 200
    assert len(calls) == 2
    # retry ใช้ key เดิม → LINE ไม่ส่งซ้ำ
    assert {c.headers["X-Line-Retry-Key"] for c in calls} == {str(key)}
    assert calls[0].headers["Authorization"] == "Bearer test-line-token"


def test_line_push_does_not_retry_client_errors():
    calls = []
    client = httpx.Client(base_url="https://api.line.me",
                          transport=httpx.MockTransport(lambda r: calls.append(r) or httpx.Response(400)))
    with pytest.raises(line_service.LineSendError):
        line_service.push_message("C", {"type": "text", "text": "x"}, uuid.uuid4(), client=client)
    assert len(calls) == 1


def test_download_token_purpose_and_binding():
    tok = line_service.create_download_token(42, 1)
    assert line_service.verify_download_token(tok) == 42


def _make_template(tmp_path, body: str):
    doc = Document()
    doc.add_paragraph(body)
    path = tmp_path / "tpl.docx"
    doc.save(path)
    return path


def test_render_docx_escapes_and_requires_all_variables(tmp_path):
    tpl = _make_template(tmp_path, "เรื่อง {{ project_title }} วงเงิน {{ budget_amount }} บาท ({{ budget_text }})")
    out = render_docx(tpl, {"project_title": "จัดซื้อ <w:p> & ครุภัณฑ์", "budget_amount": "80,000.00",
                            "budget_text": "แปดหมื่นบาทถ้วน"})
    text = "\n".join(p.text for p in Document(io.BytesIO(out)).paragraphs)
    assert "จัดซื้อ <w:p> & ครุภัณฑ์" in text          # ข้อความผู้ใช้ไม่ทำลายโครงสร้าง XML
    assert "แปดหมื่นบาทถ้วน" in text
    with pytest.raises(ExportError, match="budget_text"):
        render_docx(tpl, {"project_title": "x", "budget_amount": "1"})
