"""Integration test: วิธีเฉพาะเจาะจง ครบวงจร Login → สร้างโครงการ → workflow → Export → LINE → Dashboard → Backup/Restore.

ต้องมี PostgreSQL (TEST_DATABASE_URL) — ถ้าเชื่อมต่อไม่ได้ test นี้จะถูก skip
"""
import io
import json
import shutil
from datetime import date
from pathlib import Path

import psycopg
import pytest
from docx import Document
from sqlalchemy.engine import make_url

from app.core.config import get_settings

ROOT = Path(__file__).resolve().parents[2]
VENDOR_TAX_ID = "0105536092641"
CGD_URL = "https://www.gprocurement.go.th/new_index.html"


def _admin_conninfo() -> tuple[str, str]:
    url = make_url(str(get_settings().database_url))
    base = f"host={url.host} port={url.port or 5432} user={url.username} password={url.password}"
    return base, url.database


@pytest.fixture(scope="module")
def client():
    base, dbname = _admin_conninfo()
    try:
        with psycopg.connect(base + " dbname=postgres", autocommit=True) as c:
            c.execute(f'DROP DATABASE IF EXISTS "{dbname}" WITH (FORCE)')
            c.execute(f'CREATE DATABASE "{dbname}"')
    except psycopg.OperationalError:
        pytest.skip("PostgreSQL ไม่พร้อมใช้งาน")

    from app.core.crypto import blind_index, encrypt_field
    from app.core.security import hash_password

    with psycopg.connect(base + f" dbname={dbname}", autocommit=True) as c:
        c.execute((ROOT / "database/schema.sql").read_text())
        c.execute((ROOT / "database/seed.sql").read_text())
        c.execute("INSERT INTO departments (code, name_th) VALUES ('PSD', 'กองพัสดุ สำนักงานตัวอย่าง')")
        people = [("นาย", "สมชาย", "ใจดี", "นักวิชาการพัสดุ"), ("นาง", "สมหญิง", "รักงาน", "หัวหน้าเจ้าหน้าที่"),
                  ("นาย", "ผู้ดูแล", "ระบบ", "นักวิชาการคอมพิวเตอร์"), ("นางสาว", "ตรวจรับ", "ละเอียด", "เจ้าพนักงานพัสดุ"),
                  ("นาย", "ผู้อำนวยการ", "ตัวอย่าง", "ผู้อำนวยการ")]
        for i, (t, f, l, pos) in enumerate(people, start=1):
            cid = f"110000000000{i}"
            c.execute("""INSERT INTO personnel (department_id, title_th, first_name_th, last_name_th, position_th,
                         employment_type, citizen_id_enc, citizen_id_hmac)
                         VALUES (1, %s, %s, %s, %s, 'CIVIL_SERVANT', %s, %s)""",
                      (t, f, l, pos, encrypt_field(cid, "personnel.citizen_id"), blind_index(cid)))
        for pid, (username, role) in [(1, ("officer", "PROCUREMENT_OFFICER")), (2, ("head", "HEAD_OF_PROCUREMENT")),
                                      (3, ("admin", "ADMIN")), (5, ("director", "EXECUTIVE"))]:
            uid = c.execute("INSERT INTO users (personnel_id, username, password_hash) VALUES (%s, %s, %s) "
                            "RETURNING user_id", (pid, username, hash_password("P@ssw0rd-" + username))).fetchone()[0]
            c.execute("INSERT INTO user_roles SELECT %s, role_id FROM roles WHERE code = %s", (uid, role))

    from fastapi.testclient import TestClient
    from app.main import app
    with TestClient(app) as tc:
        yield tc


def login(client, username: str) -> dict:
    r = client.post("/api/auth/login", json={"username": username, "password": "P@ssw0rd-" + username})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def codes(resp) -> set[str]:
    return {v["code"] for v in resp.json()["violations"]}


JUSTIFICATION_56_2_B: int


def project_payload(budget: str, **kw) -> dict:
    return {"project_code": kw.get("code", "PSD-2570-001"), "fiscal_year_be": 2570,
            "title_th": "จัดซื้อวัสดุคอมพิวเตอร์ จำนวน 10 รายการ", "reason_th": "เพื่อใช้ในการปฏิบัติราชการ",
            "department_id": 1, "method_code": "SPECIFIC", "justification_id": JUSTIFICATION_56_2_B,
            "procurement_type": "GOODS", "budget_amount": budget, "delivery_days": 30,
            "items": [{"description": "หมึกพิมพ์ เลเซอร์", "quantity": "10", "unit_th": "กล่อง",
                       "unit_price": "7950.00"}]}


def test_full_specific_method_flow(client, monkeypatch):
    # --- Auth ---------------------------------------------------------------
    bad = client.post("/api/auth/login", json={"username": "officer", "password": "wrong"})
    assert bad.status_code == 401 and bad.json()["message"] == "ชื่อผู้ใช้หรือรหัสผ่านไม่ถูกต้อง"
    assert client.get("/api/methods").status_code == 401
    officer, head, admin = login(client, "officer"), login(client, "head"), login(client, "admin")
    assert client.get("/api/admin/backups", headers=officer).status_code == 403      # RBAC

    global JUSTIFICATION_56_2_B
    specific = next(m for m in client.get("/api/methods", headers=officer).json() if m["code"] == "SPECIFIC")
    JUSTIFICATION_56_2_B = next(j["id"] for j in specific["justifications"] if j["clause_ref"] == "มาตรา 56(2)(ข)")

    # --- Dynamic workflow preview: e-bidding > 5 ล้าน มีขั้นรับฟังความคิดเห็น ----------
    wf_small = client.get("/api/workflow/E_BIDDING?budget=4000000", headers=officer).json()
    wf_big = client.get("/api/workflow/E_BIDDING?budget=6000000", headers=officer).json()
    assert "PUBLIC_HEARING" not in [s["code"] for s in wf_small]
    assert "PUBLIC_HEARING" in [s["code"] for s in wf_big] and wf_big[0]["min_announce_working_days"] == 10

    # --- Project validation -------------------------------------------------
    r = client.post("/api/projects", json=project_payload("600000.00"), headers=officer)
    assert r.status_code == 422 and "BUDGET_EXCEEDS_SPECIFIC_LIMIT" in codes(r)
    r = client.post("/api/projects", json={**project_payload("80000.00"), "status": "APPROVED"}, headers=officer)
    assert r.status_code == 422                                                     # mass assignment
    r = client.post("/api/projects", json=project_payload("80000.00"), headers=officer)
    assert r.status_code == 201, r.text
    pid = r.json()["project_id"]

    # --- รายการโครงการ / หน่วยงาน ---------------------------------------------
    rows = client.get("/api/projects", params={"q": "วัสดุคอมพิวเตอร์", "fiscal_year": 2570}, headers=officer).json()
    assert [p["project_id"] for p in rows] == [pid]
    assert rows[0]["progress"] == {"done": 1, "total": 5}
    assert rows[0]["current_step"]["code"] == "PURCHASE_REQUEST_REPORT"
    assert client.get("/api/projects", params={"q": "%"}, headers=officer).json() == []      # wildcard ถูก escape
    assert client.get("/api/projects", params={"method": "E_BIDDING"}, headers=officer).json() == []
    assert client.get("/api/projects", params={"status": "HACKED"}, headers=officer).status_code == 422
    assert client.get("/api/departments", headers=officer).json()[0]["code"] == "PSD"

    def submit(step: str, data: dict, who=officer):
        version = client.get(f"/api/projects/{pid}", headers=who).json()["row_version"]
        return client.post(f"/api/projects/{pid}/steps/{step}/submit", json={"data": data, "row_version": version},
                           headers=who)

    def approve(step: str, who=head):
        return client.post(f"/api/projects/{pid}/steps/{step}/decision", json={"approve": True}, headers=who)

    # --- Workflow ordering ----------------------------------------------------
    price = {"doc_no": "พด 0002/2570", "doc_date": "2026-10-13", "vendor_tax_id": VENDOR_TAX_ID,
             "vendor_name": "บริษัท ตัวอย่างไอที จำกัด", "quoted_amount": "80000.00", "agreed_amount": "79500.00"}
    r = submit("PRICE_AGREEMENT", price)
    assert r.status_code == 422 and codes(r) == {"PREVIOUS_STEP_PENDING"}

    form = client.get(f"/api/projects/{pid}/steps/PURCHASE_REQUEST_REPORT/form", headers=officer).json()
    assert {f["name"] for f in form["fields"]} >= {"reference_price", "inspection_committee"}
    r = submit("PURCHASE_REQUEST_REPORT", {
        "doc_no": "พด 0001/2570", "doc_date": "2026-10-12", "reason_th": "วัสดุคงคลังไม่เพียงพอ",
        "spec_summary": "หมึกพิมพ์เลเซอร์ตามรุ่นที่กำหนด", "reference_price": "80000.00",
        "reference_price_source": "4", "delivery_days": 15, "evaluation_criteria": "PRICE",
        "inspection_committee": [4]})
    assert r.status_code == 200, r.text
    stale = client.post(f"/api/projects/{pid}/steps/PURCHASE_REQUEST_REPORT/submit",
                        json={"data": {}, "row_version": 1}, headers=officer)
    assert stale.status_code == 409                                                 # optimistic lock
    assert approve("PURCHASE_REQUEST_REPORT", who=officer).status_code == 403       # เจ้าหน้าที่อนุมัติไม่ได้
    assert approve("PURCHASE_REQUEST_REPORT").json()["status"] == "APPROVED"

    # --- Vendor registry ------------------------------------------------------
    r = client.post("/api/vendors", headers=officer, json={
        "tax_id": "0105536092642", "vendor_type": "JURISTIC", "name_th": "x", "address_th": "x"})
    assert r.status_code == 422                                                     # checksum ผิด
    r = client.post("/api/vendors", headers=officer, json={
        "tax_id": VENDOR_TAX_ID, "vendor_type": "JURISTIC", "name_th": "บริษัท ตัวอย่างไอที จำกัด",
        "address_th": "99 ถนนพหลโยธิน กรุงเทพมหานคร", "phone": "021234567", "bank_account": "1234567890"})
    assert r.status_code == 201, r.text
    vendor_id = r.json()["vendor_id"]

    # รายละเอียด: ข้อมูลติดต่อออกมาแบบปิดบังเท่านั้น
    v = client.get(f"/api/vendors/{vendor_id}", headers=officer).json()
    assert v["phone_masked"] == "xxx-xxx-4567" and v["bank_account_masked"] == "xxxxxx7890"
    assert "1234567890" not in str(v) and "021234567" not in str(v)
    # แก้ไข: ส่งเฉพาะช่องที่เปลี่ยน · ห้ามลบชื่อ · แก้เลขผู้เสียภาษีไม่ได้
    assert client.patch(f"/api/vendors/{vendor_id}", json={"email": "sales@example.co.th"}, headers=officer).status_code == 200
    assert client.get(f"/api/vendors/{vendor_id}", headers=officer).json()["email"] == "sales@example.co.th"
    assert client.patch(f"/api/vendors/{vendor_id}", json={"name_th": None}, headers=officer).status_code == 422
    assert client.patch(f"/api/vendors/{vendor_id}", json={"tax_id": "0105536092641"}, headers=officer).status_code == 422

    # บัญชีผู้ทิ้งงาน: เฉพาะหัวหน้าเจ้าหน้าที่ · ผู้ประกอบการที่ถูกระงับเลือกเป็นผู้ชนะไม่ได้
    other = client.post("/api/vendors", headers=officer, json={
        "tax_id": "0105536000003", "vendor_type": "JURISTIC", "name_th": "บริษัท ทิ้งงาน จำกัด", "address_th": "x"}).json()
    sanction = {"sanction_type": "ABANDONED_WORK", "reference_no": "กค (กวจ) 0405.2/ว 1", "effective_from": "2020-01-01"}
    assert client.post(f"/api/vendors/{other['vendor_id']}/sanctions", json=sanction, headers=officer).status_code == 403
    bad = {**sanction, "effective_to": "2019-01-01"}
    assert client.post(f"/api/vendors/{other['vendor_id']}/sanctions", json=bad, headers=head).status_code == 422
    assert client.post(f"/api/vendors/{other['vendor_id']}/sanctions", json=sanction, headers=head).status_code == 201
    listed = client.get("/api/vendors", params={"sanctioned": "true"}, headers=officer).json()
    assert [x["vendor_id"] for x in listed] == [other["vendor_id"]]
    assert client.get("/api/vendors", params={"q": "0-1055-36092-64-1"}, headers=officer).json()[0]["vendor_id"] == vendor_id

    assert submit("PRICE_AGREEMENT", price).status_code == 200
    assert approve("PRICE_AGREEMENT").status_code == 200

    # SoD: หัวหน้าเป็นผู้เสนอเอง แล้วอนุมัติเองไม่ได้
    award = {"doc_no": "พด 0003/2570", "doc_date": "2026-10-14", "winner_tax_id": VENDOR_TAX_ID,
             "award_amount": "79500.00"}
    assert submit("AWARD_APPROVAL", award, who=head).status_code == 200
    r = approve("AWARD_APPROVAL", who=head)
    assert r.status_code == 422 and codes(r) == {"SOD_VIOLATION"}
    assert submit("AWARD_APPROVAL", award).status_code == 422                        # ล็อกหลังส่งแล้ว

    # --- Template registry (CGD only) + Export ----------------------------------
    tpl = Document()
    tpl.add_paragraph("{{ agency_name }}")
    tpl.add_paragraph("ที่ {{ doc_no }} วันที่ {{ doc_date_th }}")
    tpl.add_paragraph("เรื่อง รายงานผลการพิจารณาและขออนุมัติสั่งซื้อ {{ project_title }}")
    tpl.add_paragraph("ผู้ได้รับคัดเลือก เลขประจำตัวผู้เสียภาษี {{ winner_tax_id }} "
                      "เป็นเงิน {{ award_amount }} บาท ({{ award_amount_text }})")
    buf = io.BytesIO()
    tpl.save(buf)
    upload = dict(document_type_code="AWARD_APPROVAL", version_label="2570.1", cgd_reference="ตัวอย่างทดสอบ",
                  effective_from=str(date.today()))
    files = {"file": ("form.docx", buf.getvalue(), "application/octet-stream")}
    r = client.post("/api/admin/templates", headers=admin, files=files,
                    data={**upload, "source_url": "https://example.com/form.docx"})
    assert r.status_code == 422                                                     # ไม่ใช่ gprocurement.go.th
    r = client.post("/api/admin/templates", headers=admin, files=files, data={**upload, "source_url": CGD_URL})
    assert r.status_code == 201, r.text
    tpl_id = r.json()["template_id"]
    assert "award_amount_text" in r.json()["variables"]["variables"]

    export_url = f"/api/projects/{pid}/steps/AWARD_APPROVAL/export"
    assert client.post(export_url, headers=officer).status_code == 422               # ยังไม่รับรองแบบฟอร์ม
    assert client.post(f"/api/admin/templates/{tpl_id}/verify", headers=head).status_code == 200

    r = client.post(export_url + "?format=docx", headers=officer)
    assert r.status_code == 200, r.text
    text = "\n".join(p.text for p in Document(io.BytesIO(r.content)).paragraphs)
    assert "79,500.00 บาท (เจ็ดหมื่นเก้าพันห้าร้อยบาทถ้วน)" in text and "14 ตุลาคม 2569" in text

    if shutil.which(get_settings().libreoffice_bin):
        r = client.post(export_url + "?format=pdf", headers=officer)
        assert r.status_code == 200 and r.content.startswith(b"%PDF")
    assert client.post(export_url + "?format=xlsx", headers=officer).status_code == 422

    director = login(client, "director")
    assert approve("AWARD_APPROVAL", who=director).json()["project_status"] == "AWARDED"

    # --- LINE share -------------------------------------------------------------
    r = client.post("/api/admin/line-channels", headers=admin,
                    json={"name_th": "กลุ่มงานพัสดุ", "line_to": "C" + "a" * 32})
    assert r.status_code == 201
    sent = {}
    from app.services import line_service

    def fake_push(to, message, retry_key, **_):
        sent.update(to=to, message=message)
        return 200
    monkeypatch.setattr(line_service, "push_message", fake_push)
    r = client.post(f"/api/projects/{pid}/share-line", headers=officer,
                    json={"channel_id": 1, "document_id": 1})
    assert r.status_code == 200, r.text
    assert sent["to"] == "C" + "a" * 32
    uri = sent["message"]["contents"]["footer"]["contents"][0]["action"]["uri"]
    token = uri.split("token=")[1]
    assert client.get(f"/api/documents/1/download?token={token}").status_code == 401   # ต้อง login ด้วย
    assert client.get(f"/api/documents/1/download?token={token}", headers=head).status_code == 200
    assert client.get(f"/api/documents/2/download?token={token}", headers=head).status_code == 403

    # --- Dashboard ------------------------------------------------------------
    d = client.get("/api/dashboard/summary?fiscal_year=2570", headers=head).json()
    assert d["kpi"]["project_count"] == 1
    assert d["kpi"]["award_total"] == "79500.00" and d["kpi"]["savings_total"] == "500.00"
    r = client.get("/api/dashboard/report.xlsx?fiscal_year=2570", headers=head)
    assert r.status_code == 200 and r.content.startswith(b"PK")

    # --- Audit chain ------------------------------------------------------------
    r = client.get("/api/admin/audit/verify", headers=admin).json()
    assert r["intact"] is True

    # --- Backup & Restore -------------------------------------------------------
    r = client.post("/api/admin/backups", headers=admin)
    assert r.status_code == 201, r.text
    backup_name = r.json()["file_name"]
    raw = (get_settings().backup_dir / backup_name).read_bytes()
    assert b"PSD-2570-001" not in raw                                               # เข้ารหัสแล้ว

    r = client.post("/api/admin/restore", headers=admin, json={
        "file_name": backup_name, "password": "wrong", "confirm_text": "ยืนยันกู้คืนข้อมูล"})
    assert r.status_code == 403
    r = client.post("/api/admin/restore", headers=admin, json={
        "file_name": "../../etc/passwd", "password": "P@ssw0rd-admin", "confirm_text": "ยืนยันกู้คืนข้อมูล"})
    assert r.status_code == 409

    # สร้างโครงการหลัง backup → restore แล้วต้องหายไป
    r = client.post("/api/projects", json=project_payload("50000.00", code="PSD-2570-002"), headers=officer)
    assert r.status_code == 422 and codes(r) == {"ITEMS_EXCEED_BUDGET"}           # รายการ 79,500 > วงเงิน
    r = client.post("/api/projects", json=project_payload("90000.00", code="PSD-2570-002"), headers=officer)
    assert r.status_code == 201, r.text
    r = client.post("/api/admin/restore", headers=admin, json={
        "file_name": backup_name, "password": "P@ssw0rd-admin", "confirm_text": "ยืนยันกู้คืนข้อมูล"})
    assert r.status_code == 200, r.text
    assert r.json()["safety_backup"].endswith("-pre-restore.dump.enc")
    officer = login(client, "officer")
    assert client.get(f"/api/projects/{pid}", headers=officer).json()["status"] == "AWARDED"
    assert client.get(f"/api/projects/{pid + 1}", headers=officer).status_code == 404
    assert client.get("/api/admin/audit/verify", headers=admin).json()["intact"] is True
    # audit ของสิ่งที่เกิดหลังเวลาสำรองต้องไม่หายไปกับการกู้คืน (ยกต่อท้าย chain ใหม่)
    assert r.json()["audit_rows_carried_over"] >= 3
    actions = [a["action"] for a in client.get("/api/admin/audit", params={"limit": 20}, headers=admin).json()["items"]]
    assert "RESTORE" in actions[:3] and {"BACKUP", "RESTORE_REQUESTED", "CREATE"} <= set(actions)
    created = client.get("/api/admin/audit", params={"action": "CREATE", "limit": 50}, headers=admin).json()["items"]
    assert any(a["entity_table"] == "projects" and a["entity_id"] == str(pid + 1) for a in created)
    manifests = client.get("/api/admin/backups", headers=admin).json()
    assert len(manifests) == 2 and json.dumps(manifests)
